"""Knowledge graph ingestion pipeline."""

import argparse
import asyncio
import hashlib
import uuid
import random
from typing import Any, Dict, List

import openpyxl
import yaml
import structlog
from sqlalchemy import select, update, text

from app.core.database import SessionLocal
from app.core.models import (
    GraphVersion, Node, Edge, NodeVisibility, NodeChunk, 
    EvalCase
)
from app.knowledge.embedders import FakeEmbedder, MistralEmbedder
from app.knowledge.validator import GraphValidator

logger = structlog.get_logger(__name__)

def make_uuid(string_id: str) -> uuid.UUID:
    return uuid.uuid5(uuid.NAMESPACE_OID, str(string_id))

def run_ingest(xlsx_path: str, domain_path: str, mapping_path: str, use_fake_embedder: bool = False) -> None:
    with open(domain_path) as f:
        domain_config = yaml.safe_load(f)
    with open(mapping_path) as f:
        mapping = yaml.safe_load(f)
        
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    
    parsed_data = {}
    for key, spec in mapping.items():
        sheet_name = spec["sheet"]
        if sheet_name not in wb.sheetnames:
            logger.warning(f"Sheet {sheet_name} not found")
            continue
        ws = wb[sheet_name]
        headers = [c.value for c in ws[1]]
        
        col_map = spec["columns"]
        col_indices = {}
        for prop, col_name in col_map.items():
            if col_name in headers:
                col_indices[prop] = headers.index(col_name)
        
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            item = {}
            for prop, idx in col_indices.items():
                item[prop] = row[idx] if row[idx] is not None else ""
            rows.append(item)
            
        parsed_data[key] = rows

    nodes = []
    edges = []
    visibilities = []
    
    def add_edge(from_id: Any, to_id: Any, type_: str, origin: str = "structural") -> None:
        if from_id and to_id:
            edges.append({"from_id": str(from_id), "to_id": str(to_id), "type": type_, "origin": origin})

    approved_conflict_articles = []
    for art in parsed_data.get("articles", []):
        nid = str(art["id"])
        status = str(art.get("status") or "").lower()
        if not status: status = "draft"
        if status == "retired": status = "archived"
        nodes.append({
            "id": nid,
            "type": "article",
            "title": str(art["title"]),
            "body": str(art["body"]),
            "status": status,
            "version": int(art.get("version") or 1),
            "owner": str(art.get("owner", "")),
            "source_ref": "workbook.xlsx",
        })
        
        vis = str(art.get("visible_roles") or "ALL")
        for role in (r.strip() for r in vis.split(",")):
            visibilities.append({"node_id": nid, "role_id": role})
            
        if art.get("supersedes"):
            add_edge(nid, art["supersedes"], "supersedes")
            
        if status == "approved" and art.get("conflict_key"):
            approved_conflict_articles.append((nid, art["conflict_key"]))
            
    for i, (id1, key1) in enumerate(approved_conflict_articles):
        for j, (id2, key2) in enumerate(approved_conflict_articles):
            if i < j and key1 == key2:
                add_edge(id1, id2, "conflicts_with", "computed")
                add_edge(id2, id1, "conflicts_with", "computed")

    for wf in parsed_data.get("workflows", []):
        nid = str(wf["id"])
        nodes.append({
            "id": nid,
            "type": "workflow",
            "title": str(wf["title"]),
            "body": str(wf.get("description", "")),
            "status": "approved",
            "version": 1,
            "owner": str(wf.get("owner_team", "")),
            "source_ref": "workbook.xlsx",
        })
        visibilities.append({"node_id": nid, "role_id": "ALL"})
        add_edge(nid, wf.get("article_id"), "defined_in")
        add_edge(nid, wf.get("owner_team"), "owned_by")
        add_edge(nid, wf.get("escalation_team"), "escalates_to")
        add_edge(nid, wf.get("system"), "performed_in")
        if wf.get("form_required"):
            add_edge(nid, wf.get("form_required"), "requires_form")
            
    for step in parsed_data.get("steps", []):
        wf_id = str(step["workflow_id"])
        step_order = step["step_order"]
        nid = f"{wf_id}_Step_{step_order}"
        nodes.append({
            "id": nid,
            "type": "step",
            "title": str(step["title"]),
            "body": str(step.get("description", "")),
            "status": "approved",
            "version": 1,
            "owner": str(step.get("team", "")),
            "source_ref": "workbook.xlsx",
        })
        visibilities.append({"node_id": nid, "role_id": "ALL"})
        add_edge(wf_id, nid, "has_step")
        add_edge(nid, step.get("team"), "performed_by")
        if step.get("system"):
            add_edge(nid, step.get("system"), "performed_in")
        if step.get("form"):
            add_edge(nid, step.get("form"), "requires_form")
            
    extra_nodes = set()
    node_types = {} 
    def add_extra(nid: Any, ntype: str) -> None:
        if nid and str(nid).strip() and nid not in node_types:
            node_types[nid] = ntype
            extra_nodes.add((nid, ntype))
            
    for wf in parsed_data.get("workflows", []):
        add_extra(wf.get("owner_team"), "team")
        add_extra(wf.get("escalation_team"), "team")
        add_extra(wf.get("system"), "system")
        add_extra(wf.get("form_required"), "form")
    for step in parsed_data.get("steps", []):
        add_extra(step.get("team"), "team")
        add_extra(step.get("system"), "system")
        add_extra(step.get("form"), "form")
    for f in parsed_data.get("field_definitions", []):
        add_extra(f.get("name"), "field")
        
    for nid, ntype in extra_nodes:
        nodes.append({
            "id": str(nid), "type": ntype, "title": str(nid), "body": "", "status": "approved", "version": 1, "owner": "", "source_ref": "workbook.xlsx"
        })
        visibilities.append({"node_id": str(nid), "role_id": "ALL"})
        
    for sf in parsed_data.get("step_fields", []):
        wf_id = str(sf["workflow_id"])
        order = sf["step_order"]
        step_nid = f"{wf_id}_Step_{order}"
        add_edge(step_nid, sf["field_name"], "requires")
        
    val = GraphValidator(domain_config)
    errors = val.validate(nodes, edges)
    
    if errors:
        logger.error("Validation failed", errors=errors)
        raise ValueError(f"Graph validation failed: {errors}")
        
    logger.info("Graph validated successfully", nodes=len(nodes), edges=len(edges))

    version_id = uuid.uuid4()
    
    with SessionLocal() as session:
        # Load Eval Cases
        ops = parsed_data.get("operations_requests", [])
        if ops:
            random.shuffle(ops)
            split_idx = int(len(ops) * 0.3)
            seed_ops = ops[:split_idx]
            heldout_ops = ops[split_idx:]
            
            session.execute(text("TRUNCATE eval_cases CASCADE;"))
            
            for split_name, ops_list in [("seed", seed_ops), ("heldout", heldout_ops)]:
                for op in ops_list:
                    ec = EvalCase(
                        split=split_name,
                        input_json={"request_text": op["request_text"], "intent": op.get("intent", "")},
                        expected_json={
                            "correct_team": op["correct_team"],
                            "correct_outcome": op["correct_outcome"],
                            "correct_workflow": op["correct_workflow"]
                        }
                    )
                    session.add(ec)
            
        draft = GraphVersion(id=version_id, status="draft")
        session.add(draft)
        
        embedder = FakeEmbedder() if use_fake_embedder else MistralEmbedder()
        
        res = session.execute(select(NodeChunk.content_hash))
        existing_hashes = {r[0] for r in res.all()}
        
        node_db_objs = []
        chunks_to_insert = []
        texts_to_embed = []
        
        for n in nodes:
            db_id = make_uuid(n["id"])
            content_hash = hashlib.sha256((n["title"] + n["body"]).encode()).hexdigest()
            node_db = Node(
                id=db_id,
                version_id=version_id,
                type=n["type"],
                title=n["title"],
                body=n["body"],
                status=n["status"],
                version=n["version"],
                owner=n["owner"],
                source_ref=n["source_ref"],
                content_hash=content_hash,
            )
            node_db_objs.append(node_db)
            
            text_chunk = f"{n['title']}\n\n{n['body']}"
            chunk = NodeChunk(
                id=uuid.uuid4(),
                node_id=db_id,
                version_id=version_id,
                text=text_chunk,
                content_hash=content_hash,
                embedding_model="mistral" if not use_fake_embedder else "fake"
            )
            chunks_to_insert.append(chunk)
            
            if content_hash not in existing_hashes:
                texts_to_embed.append((chunk, text_chunk))
        
        session.add_all(node_db_objs)
        
        if texts_to_embed:
            logger.info(f"Embedding {len(texts_to_embed)} new chunks")
            texts = [t for _, t in texts_to_embed]
            
            # Use asyncio.run for calling the async embedder in synchronous context
            import asyncio
            try:
                loop = asyncio.get_running_loop()
            except RuntimeError:
                embeddings = asyncio.run(embedder.embed(texts))
            else:
                raise RuntimeError("run_ingest called from async loop, use sync embedder or thread")
                
            for (chunk, _), emb in zip(texts_to_embed, embeddings):
                chunk.embedding = emb
                
        if len(texts_to_embed) < len(chunks_to_insert):
            needed_hashes = [c.content_hash for c in chunks_to_insert if c.embedding is None]
            if needed_hashes:
                res = session.execute(
                    select(NodeChunk).where(NodeChunk.content_hash.in_(needed_hashes))
                )
                hash_to_emb = {}
                for r in res.scalars():
                    if r.embedding is not None and r.content_hash not in hash_to_emb:
                        hash_to_emb[r.content_hash] = r.embedding
                
                for c in chunks_to_insert:
                    if c.embedding is None:
                        c.embedding = hash_to_emb.get(c.content_hash)
                        
        session.add_all(chunks_to_insert)
        
        for e in edges:
            session.add(Edge(
                from_id=make_uuid(e["from_id"]),
                to_id=make_uuid(e["to_id"]),
                type=e["type"],
                origin=e["origin"],
                version_id=version_id
            ))
            
        for v in visibilities:
            session.add(NodeVisibility(
                node_id=make_uuid(v["node_id"]),
                version_id=version_id,
                role_id=v["role_id"]
            ))
            
        session.execute(
            update(GraphVersion).where(GraphVersion.status == "active").values(status="archived")
        )
        draft.status = "active"
        
        session.commit()
        logger.info("Ingestion complete", version_id=str(version_id))

if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx_path")
    parser.add_argument("--domain", default="/app/data/domain.yaml")
    parser.add_argument("--mapping", default="/app/data/mapping.yaml")
    parser.add_argument("--fake-embedder", action="store_true")
    args = parser.parse_args()
    
    run_ingest(args.xlsx_path, args.domain, args.mapping, args.fake_embedder)
