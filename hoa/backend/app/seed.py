"""Database seeder for the hackathon project."""

import asyncio
import json
import logging
import os
from pathlib import Path

import openpyxl
import yaml
from sqlalchemy import text

from app.core.database import SessionLocal
from app.core.models import (
    DomainConfig,
    FieldDefinition,
    RoutingRule,
    PolicyThreshold,
)
from app.knowledge.ingest import run_ingest
from app.core.user_seeder import seed_users

logger = logging.getLogger(__name__)

def seed_db() -> None:
    print("Generating synthetic data...")
    os.system("python data/synthetic_generator.py")
    
    print("Running knowledge graph ingest...")
    run_ingest(
        xlsx_path="data/workbook.xlsx",
        domain_path="data/domain.yaml",
        mapping_path="data/mapping.yaml",
        use_fake_embedder=True
    )
    
    print("Seeding domain configs, rules and thresholds...")
    with open("data/domain.yaml") as f:
        domain = yaml.safe_load(f)
        
    wb = openpyxl.load_workbook("data/workbook.xlsx", data_only=True)
    
    with SessionLocal() as session:
        session.execute(text("TRUNCATE domain_config, field_definitions, routing_rules, policy_thresholds CASCADE;"))
        
        session.add(DomainConfig(key="roles", value={"roles": domain.get("roles", [])}))
        session.add(DomainConfig(key="categories", value={"categories": domain.get("categories", [])}))
        session.add(DomainConfig(key="intents", value={"intents": domain.get("intents", [])}))
        session.add(DomainConfig(key="prompt_injection_patterns", value={"patterns": domain.get("prompt_injection_patterns", [])}))
        
        if "routing_rules" in wb.sheetnames:
            ws = wb["routing_rules"]
            headers = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                item = dict(zip(headers, row))
                rule = RoutingRule(
                    intent=item["intent"],
                    team=item["route_to"],
                    urgency=item.get("urgency", "normal"),
                    condition_json=json.loads(item["condition_json"]) if item.get("condition_json") else None
                )
                session.add(rule)
                
        if "policy_thresholds" in wb.sheetnames:
            ws = wb["policy_thresholds"]
            headers = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                item = dict(zip(headers, row))
                pt = PolicyThreshold(
                    key=item["policy_key"],
                    value=float(item["threshold_value"]),
                    description=item.get("description", "")
                )
                session.add(pt)
                
        if "field_definitions" in wb.sheetnames:
            ws = wb["field_definitions"]
            headers = [c.value for c in ws[1]]
            for row in ws.iter_rows(min_row=2, values_only=True):
                if not row[0]: continue
                item = dict(zip(headers, row))
                fd = FieldDefinition(
                    workflow_key=item.get("workflow_id", ""),
                    field_name=item["name"],
                    field_type=item.get("type", "string"),
                    required=item.get("required", True),
                    validation_regex=item.get("validation_regex")
                )
                session.add(fd)
                
        session.commit()
    
    print("Seeding demo users...")
    with SessionLocal() as session:
        seed_users(session)
    
    print("Seeding complete.")

if __name__ == "__main__":
    seed_db()
